"""评课课件、大纲和知识图谱文件的最小文本预处理。"""

from __future__ import annotations

import csv
import io
import json
import logging
from pathlib import Path
from typing import Any

import aiohttp

from media_platform.contracts.content_evaluation import FileMetadata, FileType
from services.content_analysis.infrastructure.clients import safe_url


LOGGER = logging.getLogger(__name__)


class EvaluationFilePreprocessor:
    def __init__(self, *, timeout_seconds: float = 60, max_size_mb: int = 100):
        self.timeout_seconds = timeout_seconds
        self.max_size_bytes = max_size_mb * 1024 * 1024

    async def _download(self, file: FileMetadata) -> bytes:
        url = str(file.file_url)
        LOGGER.info("下载评课材料: file=%s, url=%s", file.file_name, safe_url(url))
        timeout = aiohttp.ClientTimeout(total=self.timeout_seconds)
        async with aiohttp.ClientSession(timeout=timeout) as session:
            async with session.get(url) as response:
                response.raise_for_status()
                data = await response.read()
        if len(data) > self.max_size_bytes:
            raise RuntimeError(f"评课材料超过大小限制: {file.file_name}")
        return data

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise RuntimeError("材料文本编码无法识别")

    def _extract(self, file: FileMetadata, data: bytes, *, extract_images: bool) -> str:
        file_type = file.file_type
        suffix = Path(file.file_name).suffix.lower()
        if file_type == FileType.TXT:
            return self._decode_text(data)
        if file_type == FileType.JSON:
            return json.dumps(json.loads(self._decode_text(data)), ensure_ascii=False, indent=2)
        if file_type == FileType.CSV:
            rows = csv.reader(io.StringIO(self._decode_text(data)))
            return "\n".join("\t".join(row) for row in rows)
        if file_type == FileType.HTML:
            from bs4 import BeautifulSoup
            return BeautifulSoup(self._decode_text(data), "html.parser").get_text("\n", strip=True)
        if file_type == FileType.PDF:
            import fitz
            document = fitz.open(stream=data, filetype="pdf")
            try:
                return "\n".join(page.get_text("text") for page in document)
            finally:
                document.close()
        if file_type == FileType.DOCX:
            from docx import Document
            document = Document(io.BytesIO(data))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        if file_type == FileType.PPT or suffix in {".ppt", ".pptx"}:
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(data))
            texts: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
            return "\n".join(texts)
        if file_type == FileType.IMAGE and extract_images:
            from PIL import Image
            import pytesseract
            return pytesseract.image_to_string(Image.open(io.BytesIO(data)), lang="chi_sim+eng")
        raise RuntimeError(f"不支持的评课材料类型: {file_type.value}")

    async def process(self, file: FileMetadata, *, category: str, extract_images: bool) -> dict[str, Any]:
        data = await self._download(file)
        text = self._extract(file, data, extract_images=extract_images).strip()
        if not text:
            raise RuntimeError(f"评课材料未提取到文本: {file.file_name}")
        return {
            "category": category,
            "file_name": file.file_name,
            "file_type": file.file_type.value,
            "text": text,
        }


__all__ = ["EvaluationFilePreprocessor"]
